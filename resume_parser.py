import os
import re
from typing import Dict, List, Tuple, Any

SKILL_LEXICON = {
    "Python": ["python", "python3", "py", "pandas", "numpy", "django", "flask", "fastapi", "scikit-learn", "pytorch", "tensorflow"],
    "Java": ["java", "spring", "springboot", "spring boot", "hibernate", "maven", "jvm", "j2ee"],
    "C/C++": ["c++", "c/c++", "cpp", "c programming", "embedded c", "pointers", "stl"],
    "Web Development": ["web development", "html", "html5", "css", "css3", "javascript", "typescript", "react", "node.js", "nodejs", "vue", "angular", "bootstrap", "tailwind", "rest api"],
    "Data Analytics": ["data analytics", "data analysis", "eda", "exploratory data analysis", "statistics", "business analytics", "kpis", "reporting", "data cleaning", "etl"],
    "AI/ML": ["ai/ml", "machine learning", "artificial intelligence", "deep learning", "nlp", "computer vision", "neural networks", "regression", "classification", "scikit-learn", "llm", "genai"],
    "SQL": ["sql", "postgresql", "mysql", "sqlite", "oracle", "sql server", "rdbms", "queries", "joins", "database design", "stored procedures"],
    "Cloud": ["cloud", "aws", "amazon web services", "azure", "google cloud", "gcp", "docker", "kubernetes", "devops", "s3", "ec2", "lambda"],
    "Excel": ["excel", "ms excel", "microsoft excel", "vlookup", "xlookup", "pivot tables", "spreadsheets", "macros", "formulas"],
    "Data Visualization": ["data visualization", "power bi", "powerbi", "tableau", "matplotlib", "seaborn", "dashboards", "plotly", "looker"],
    "Communication": ["communication", "presentation", "public speaking", "written communication", "technical writing", "client communication", "stakeholder management"],
    "Teamwork": ["teamwork", "collaboration", "cross-functional", "peer review", "team player", "scrum", "agile"],
    "Problem Solving": ["problem solving", "analytical thinking", "debugging", "algorithm design", "critical thinking", "root cause analysis"],
    "Leadership": ["leadership", "team lead", "lead", "mentored", "initiative", "managed team", "organized", "coordinator"],
    "Time Management": ["time management", "deadline", "prioritization", "task management", "multitasking", "sprint planning"],
}

def extract_text_from_file(file_path: str) -> str:
    """Extract text from PDF, DOCX, or TXT file."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    pages_text.append(t)
            text = "\n".join(pages_text)
        except Exception as e:
            print("pypdf extraction error:", e)
            try:
                with open(file_path, "rb") as f:
                    content = f.read().decode("latin-1", errors="ignore")
                    words = re.findall(r"[A-Za-z0-9+#./-]{2,}", content)
                    text = " ".join(words)
            except Exception:
                text = ""
                
    elif ext in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            paragraphs.append(cell.text)
            text = "\n".join(paragraphs)
        except Exception as e:
            print("docx extraction error:", e)
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(file_path) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    text = "".join(node.text for node in tree.iter() if node.text)
            except Exception:
                text = ""
                
    elif ext in (".txt", ".md", ".rtf"):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception:
            with open(file_path, "r", encoding="latin-1", errors="ignore") as f:
                text = f.read()
                
    return text


def parse_resume_text(text: str, current_scores: Dict[str, float] = None, target_role_required: Dict[str, float] = None) -> Dict[str, Any]:
    """
    Parse resume text, identify skills with confidence scores,
    and compare against current profile and target career goal.
    """
    current_scores = current_scores or {}
    target_role_required = target_role_required or {}
    lower_text = " " + text.lower() + " "
    
    detected_skills = {}
    matched_keywords = {}
    
    # Project & Experience context detection
    has_projects_section = bool(re.search(r"(projects?|portfolio|personal projects|academic projects)", lower_text))
    has_experience_section = bool(re.search(r"(experience|internship|work experience|employment)", lower_text))
    
    for canonical, keywords in SKILL_LEXICON.items():
        count = 0
        found_words = []
        for kw in keywords:
            pattern = r"(?:\b|_)" + re.escape(kw) + r"(?:\b|_)"
            matches = re.findall(pattern, lower_text)
            if matches:
                count += len(matches)
                found_words.append(kw)
                
        if count > 0:
            base = 60
            if count >= 2:
                base += 10
            if count >= 4:
                base += 10
            if has_projects_section:
                base += 5
            if has_experience_section and count >= 2:
                base += 5
            suggested_score = min(90, max(50, base))
            detected_skills[canonical] = suggested_score
            matched_keywords[canonical] = list(set(found_words))
            
    comparison = []
    all_skills = sorted(list(set(list(SKILL_LEXICON.keys()) + list(current_scores.keys()))))
    
    new_skills = []
    improved_skills = []
    missing_for_target = []
    
    for skill in all_skills:
        curr = current_scores.get(skill, 0)
        found = detected_skills.get(skill)
        req = target_role_required.get(skill, 0)
        
        status = "not_found"
        if found is not None:
            if curr == 0:
                status = "new"
                new_skills.append({"skill": skill, "suggested_score": found, "keywords": matched_keywords.get(skill, [])})
            elif found > curr:
                status = "higher"
                improved_skills.append({"skill": skill, "current_score": curr, "suggested_score": found})
            else:
                status = "confirmed"
        else:
            if req > 0 and curr < req:
                missing_for_target.append({"skill": skill, "required_score": req, "current_score": curr})
                
        comparison.append({
            "skill": skill,
            "current_score": curr,
            "detected_in_resume": found is not None,
            "suggested_score": found if found is not None else curr,
            "keywords": matched_keywords.get(skill, []),
            "target_required": req,
            "status": status,
        })
        
    return {
        "text_length": len(text),
        "detected_count": len(detected_skills),
        "detected_skills": detected_skills,
        "matched_keywords": matched_keywords,
        "new_skills": new_skills,
        "improved_skills": improved_skills,
        "missing_for_target": missing_for_target,
        "comparison": comparison,
    }
